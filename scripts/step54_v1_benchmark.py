"""STEP54 fixed V1 benchmark harness.  It does not modify production code."""
from __future__ import annotations

import csv
import json
import subprocess
import sys
from pathlib import Path
from statistics import median
from time import perf_counter
from typing import Any

import step52_blind_test as product

ROOT = Path(__file__).resolve().parents[1] / "artifacts" / "step54_v1_benchmark"
SITES = [
    ("NAURA", "https://career.naura.com/campus/jobs"),
    ("AECC", "https://aecc.iguopin.com/job"),
    ("Sinomach", "https://zhaopin.sinomach.com.cn/SU64b4cfe82f9d24760ae8b80c/pb/school.html"),
    ("AMEC", "https://app.mokahr.com/campus_apply/amec/4362#/jobs?page=1&anchorName=jobsList"),
    ("YMTC", "https://ymtc.zhiye.com/campus/jobs"),
    ("CXMT", "https://cxmt.zhiye.com/campus/jobs"),
    ("CATL", "https://app.mokahr.com/campus-recruitment/catlhr/148948#/jobs"),
    ("Guangzhou Metro", "https://gzmetro.zhiye.com/campus/jobs"),
    ("Geely", "https://careers.geelytech.com/campus"),
    ("Hisense", "https://jobs.hisense.com/campus/jobs"),
]


def slug(name: str) -> str:
    return product._slug(name)


def load(path: Path, fallback: Any) -> Any:
    try:
        return json.loads(path.read_text(encoding="utf-8"))
    except Exception:
        return fallback


def env_failure(reason: str) -> bool:
    text = reason.lower()
    return any(token in text for token in ("dns", "network", "browserruntime", "browser_launch", "navigation_timeout", "targetclosed", "connection", "name or service"))


def excel_check(path: Path) -> tuple[bool, int, str]:
    if not path.exists() or path.stat().st_size == 0:
        return False, 0, "MISSING_OR_EMPTY"
    try:
        from openpyxl import load_workbook
        workbook = load_workbook(path, read_only=True, data_only=True)
        rows = sum(max(0, sheet.max_row - 1) for sheet in workbook.worksheets)
        workbook.close()
        return True, rows, "OPENED"
    except Exception as exc:
        return False, 0, type(exc).__name__


def classify(row: dict[str, Any]) -> str:
    reason = str(row.get("failure_reason") or "")
    if row.get("collection_status") == "COMPLETE" and row.get("false_complete_guard_triggered"):
        return "IDENTITY_FAILURE"
    if env_failure(reason):
        return "ENVIRONMENT_BROWSER" if "browser" in reason.lower() or "targetclosed" in reason.lower() else "ENVIRONMENT_NETWORK"
    if not row.get("source_found"):
        return "SOURCE_DISCOVERY_FAILURE"
    if not row.get("plan_executable") or not row.get("plan_valid") or not row.get("dispatchable"):
        return "PLAN_NOT_EXECUTABLE"
    if row.get("collection_status") != "COMPLETE":
        if "page" in reason.lower() or "pagination" in reason.lower() or row.get("stagnant_pages"):
            return "PAGINATION_FAILURE"
        if row.get("reported_total") is not None and row.get("unique_jobs", 0) < row["reported_total"]:
            return "TOTAL_MISMATCH"
        return "PARTIAL_COLLECTION"
    if row.get("jd_incomplete", 0):
        return "JD_COMPLETENESS_FAILURE"
    if not row.get("excel_success"):
        return "EXPORT_FAILURE"
    return "SUCCESS"


def enrich(base: dict[str, Any], name: str, url: str) -> dict[str, Any]:
    folder = ROOT / "per_site" / slug(name)
    discovery = load(folder / "discovery.json", {})
    plan = load(folder / "plan.json", {})
    collection = load(folder / "collection.json", {})
    metrics = collection.get("metrics") or {}
    complete = collection.get("data_completeness") or {}
    audit = collection.get("duplicate_audit") or {}
    excel_ok, excel_rows, excel_state = excel_check(folder / "reports" / "jobs.xlsx")
    total = collection.get("total_expected")
    raw = int(collection.get("total_fetched") or metrics.get("raw_rows") or 0)
    unique = int(collection.get("total_unique") or 0)
    jd_total = int(complete.get("total_jobs") or 0)
    jd_complete = int(complete.get("jd_complete") or 0)
    source_type = "RUNTIME_STATE" if discovery.get("runtime_source") else ("NETWORK_JSON" if discovery.get("probable_list_api") else "DOM_LIST" if (discovery.get("dom_fallback") or {}).get("status") else "NONE")
    source_false = bool(base.get("source_false_positive"))
    false_complete = bool(collection.get("status") == "COMPLETE" and isinstance(total, int) and unique != total)
    row = {
        "site_name": name, "url": url, "input_valid": True,
        "environment_status": "EVALUABLE", "provider_fingerprint": base.get("provider"),
        "discovery_status": discovery.get("status") or "NOT_RUN", "source_found": bool(base.get("source_found")),
        "source_type": source_type, "source_confidence": base.get("source_confidence"),
        "plan_mode": plan.get("mode"), "plan_executable": bool(base.get("plan_executable")),
        "plan_valid": bool(base.get("plan_valid")), "dispatchable": bool(base.get("dispatchable")),
        "pagination_mode": plan.get("pagination_type"), "reported_total": total,
        "pages_requested": int(metrics.get("pages_requested") or metrics.get("list_pages") or 0),
        "pages_succeeded": int(metrics.get("pages_succeeded") or metrics.get("list_pages") or 0),
        "raw_jobs": raw, "unique_jobs": unique, "duplicate_jobs": max(0, raw - unique),
        "collection_status": collection.get("status") or base.get("collection_status"),
        "termination_reason": metrics.get("termination_reason") or base.get("failure_reason") or "",
        "jd_complete": jd_complete, "jd_incomplete": max(0, jd_total - jd_complete),
        "jd_complete_rate": round(jd_complete / jd_total, 6) if jd_total else 0.0,
        "detail_required": int(complete.get("detail_required") or 0), "detail_attempted": int(complete.get("detail_attempted") or 0),
        "detail_succeeded": int(complete.get("detail_succeeded") or 0), "detail_failed": int(complete.get("detail_failed") or 0),
        "scope_mismatch": False, "false_complete_guard_triggered": false_complete,
        "excel_success": excel_ok, "excel_rows": excel_rows, "excel_validation": excel_state,
        "json_success": (folder / "reports" / "jobs.json").exists(), "markdown_success": (folder / "reports" / "report.md").exists(),
        "elapsed_seconds": float(base.get("elapsed_seconds") or 0), "discovery_seconds": float(discovery.get("elapsed_seconds") or 0),
        "collection_seconds": float(metrics.get("elapsed_seconds") or 0), "export_seconds": float(metrics.get("export_seconds") or 0),
        "source_false_positive": source_false, "artifact_path": str(folder), "original_reason": base.get("failure_reason") or "",
    }
    row["final_failure_class"] = classify(row)
    if row["final_failure_class"].startswith("ENVIRONMENT_"):
        row["environment_status"] = row["final_failure_class"]
    row["e2e_success"] = row["final_failure_class"] == "SUCCESS" and not source_false and not false_complete and not row["scope_mismatch"]
    (folder / "benchmark_row.json").write_text(json.dumps(row, ensure_ascii=False, indent=2), encoding="utf-8")
    return row


def rate(rows: list[dict[str, Any]], fn) -> float:
    return round(sum(bool(fn(row)) for row in rows) / len(rows), 4) if rows else 0.0


def metrics(rows: list[dict[str, Any]]) -> dict[str, float]:
    jd_denominator = sum(x["jd_complete"] + x["jd_incomplete"] for x in rows)
    return {"source_found_rate": rate(rows, lambda x: x["source_found"] and not x["source_false_positive"]),
            "executable_plan_rate": rate(rows, lambda x: x["plan_executable"] and x["plan_valid"] and x["dispatchable"]),
            "collection_started_rate": rate(rows, lambda x: x["collection_status"] not in (None, "NOT_STARTED")),
            "full_collection_rate": rate(rows, lambda x: x["collection_status"] == "COMPLETE" and not x["false_complete_guard_triggered"]),
            "jd_complete_rate": round(sum(x["jd_complete"] for x in rows) / jd_denominator, 4) if jd_denominator else 0.0,
            "excel_success_rate": rate(rows, lambda x: x["excel_success"]), "e2e_success_rate": rate(rows, lambda x: x["e2e_success"])}


def write_outputs(rows: list[dict[str, Any]], policy: dict[str, Any]) -> None:
    fields = list(rows[0])
    for name, selection in (("benchmark_results.csv", rows), ("failure_matrix.csv", rows), ("jd_metrics.csv", rows), ("export_metrics.csv", rows)):
        with (ROOT / name).open("w", encoding="utf-8", newline="") as output:
            writer = csv.DictWriter(output, fieldnames=fields); writer.writeheader(); writer.writerows(selection)
    before = {"source_found_rate": .2, "executable_plan_rate": .2, "collection_started_rate": .2, "full_collection_rate": 0.0, "jd_complete_rate": 0.0, "excel_success_rate": 0.0, "e2e_success_rate": 0.0}
    all_metrics = metrics(rows); evaluable = [x for x in rows if x["environment_status"] == "EVALUABLE"]
    after_evaluable = metrics(evaluable)
    comparisons = [{"metric": key, "step52d": before[key], "step54_all_valid": all_metrics[key], "absolute_change": round(all_metrics[key]-before[key],4), "percentage_point_change": round((all_metrics[key]-before[key])*100,2)} for key in before]
    with (ROOT / "benchmark_before_after.csv").open("w", encoding="utf-8", newline="") as output:
        writer = csv.DictWriter(output, fieldnames=list(comparisons[0])); writer.writeheader(); writer.writerows(comparisons)
    elapsed = [x["elapsed_seconds"] for x in rows]
    summary = {"benchmark_id": "STEP54_V1_FIXED_10", "policy": policy, "all_valid_inputs": all_metrics,
               "algorithmically_evaluable": {"denominator": len(evaluable), **after_evaluable}, "before_after": comparisons,
               "source_false_positive_count": sum(x["source_false_positive"] for x in rows),
               "false_complete_count": sum(x["false_complete_guard_triggered"] for x in rows),
               "performance": {"avg": round(sum(elapsed)/len(elapsed),3), "median": round(median(elapsed),3), "max": max(elapsed), "slowest": max(rows,key=lambda x:x["elapsed_seconds"])["site_name"]}, "sites": rows}
    (ROOT / "benchmark_summary.json").write_text(json.dumps(summary, ensure_ascii=False, indent=2), encoding="utf-8")


def main() -> None:
    ROOT.mkdir(parents=True, exist_ok=True); (ROOT / "per_site").mkdir(exist_ok=True)
    status = subprocess.run(["git", "status", "--short"], capture_output=True, text=True, check=False).stdout.splitlines()
    commit = subprocess.run(["git", "rev-parse", "HEAD"], capture_output=True, text=True, check=False).stdout.strip()
    policy = {"benchmark_id": "STEP54_V1_FIXED_10", "commit": commit, "working_tree_status": status,
              "browser": "production BrowserRuntime; fresh run per site", "discovery_budget_seconds": 60,
              "discovery_timeout_ms": 15000, "collection_budget_seconds": 100, "http_max_pages": 1000,
              "http_max_retries": 1, "export": "ReportManager + openpyxl workbook-open validation"}
    (ROOT / "benchmark_policy.json").write_text(json.dumps(policy, ensure_ascii=False, indent=2), encoding="utf-8")
    product.ROOT = ROOT / "per_site"; product.DISCOVERY_BUDGET = 60; product.DISCOVERY_TIMEOUT_MS = 15000; product.COLLECTION_BUDGET = 100
    rows = []
    for name, url in SITES:
        base = product._run_site((name, "Fixed V1 benchmark", url))
        rows.append(enrich(base, name, url))
        print(f"STEP54 {name}: {rows[-1]['final_failure_class']} e2e={rows[-1]['e2e_success']}", flush=True)
    write_outputs(rows, policy)


if __name__ == "__main__":
    if "--render-existing" in sys.argv:
        rows = [load(ROOT / "per_site" / slug(name) / "benchmark_row.json", {}) for name, _url in SITES]
        write_outputs(rows, load(ROOT / "benchmark_policy.json", {}))
    else:
        main()
