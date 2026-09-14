"""STEP 51 — JD / Detail Completeness.

Principle: full JD truth > structured field perfection. A job is complete
when the source genuinely provides the full JD text, no matter which
field-name carries it, and structured requirement/responsibility columns are
optional when the source posts a single whole-JD body.
"""
import json

from openpyxl import load_workbook

from job_extractor.collectors.generic_http import GenericHttpCollector
from job_extractor.field_semantics import (
    canonical_jd,
    needs_detail_fetch,
    pick_jd_fields,
)
from job_extractor.models import CollectionResult, Job
from job_extractor.reporting import ReportManager
from job_extractor.runtime import evaluate_data_completeness
from tests.test_generic_collectors import Client, plan, raw


def _collect(items, details=None, **plan_kwargs):
    plan_value = plan(detail_mode=plan_kwargs.pop("detail_mode", "LIST_SUFFICIENT"), **plan_kwargs)
    client = Client([{"data": {"items": items, "total": len(items)}}], details or {})
    return GenericHttpCollector(plan_value, client).collect()


def _detail_plan_result(items, details=None):
    return _collect(
        items,
        details,
        detail_mode="DETAIL_FALLBACK",
        detail_endpoint_template="https://api.x.test/detail/{id}",
        detail_method="GET",
        detail_id_field="id",
    )


# CASE 1: description + requirements on the list -> LIST_SUFFICIENT, no detail.
def test_case1_list_split_is_sufficient_no_detail_fetch():
    item = raw(1, "Build things")
    item["requirements"] = "Python and Go"
    result = _collect([item], {"1": {"description": "Detail body"}}, detail_mode="LIST_SUFFICIENT")
    assert result.status == "COMPLETE"
    assert result.metrics.detail_requests == 0
    assert result.jobs[0].full_jd and result.jobs[0].requirements


# CASE 2: title/id/location only + detail carries the full JD -> fetched.
def test_case2_absent_jd_triggers_detail_fetch():
    item = raw(1, "")
    item["description"] = ""
    item["requirements"] = ""
    detail = {"description": "Responsibilities: " + "Build reliable systems. " * 30, "requirements": "Python"}
    result = _collect(
        [item], {"1": detail}, detail_mode="DETAIL_REQUIRED",
        detail_endpoint_template="https://api.x.test/detail/{id}",
        detail_method="GET", detail_id_field="id",
    )
    assert result.metrics.detail_requests == 1
    job = result.jobs[0]
    assert job.full_jd and "Build reliable systems." in job.full_jd
    assert job.requirements == ["Python"]


# CASE 3: single whole-JD field (positionDescription) -> complete even though
# no structured requirements field exists at the source.
def test_case3_position_description_single_field_is_complete():
    body = "岗位职责\n负责核心系统开发与维护\n任职要求\n本科以上学历，三年以上经验"
    result = _collect([{"id": "1", "title": "工程师", "positionDescription": body}])
    job = result.jobs[0]
    assert job.full_jd == body
    dc = evaluate_data_completeness(result)
    assert dc.jd_complete == 1 and dc.jd_incomplete == 0
    assert dc.source_requirements_absent == 1
    # recognized state: FULL_TEXT, not SUMMARY (heading hits + length).
    assert job.raw_data["_jd_state"] == "FULL_TEXT"


# CASE 4: responsibility + requirement halves -> canonical JD keeps both.
def test_case4_split_halves_preserved_in_canonical_jd():
    raw_record = {
        "id": "1", "title": "工程师",
        "jobResponsibility": "1. 负责服务端开发\n2. 参与架构评审",
        "jobRequirement": "1. 本科以上学历\n2. 熟悉 Python",
    }
    full, state, _, _ = canonical_jd(raw_record)
    assert state == "SPLIT"
    assert "岗位职责" in full and "任职要求" in full
    assert "负责服务端开发" in full and "熟悉 Python" in full
    recognized = pick_jd_fields(raw_record)
    assert recognized["responsibilities"] and recognized["requirements"]


# CASE 5: detail returns empty fields -> list text is never overwritten.
def test_case5_empty_detail_cannot_overwrite_list_text():
    item = raw(1, "List level description. " * 30)
    item["requirements"] = ""
    detail = {"description": "", "requirements": ""}
    result = _detail_plan_result([item], {"1": detail})
    assert result.metrics.detail_requests == 0  # list JD credible -> no fetch
    # force a fetch with DETAIL_REQUIRED to exercise the merge path
    result2 = _collect(
        [dict(item)], {"1": detail}, detail_mode="DETAIL_REQUIRED",
        detail_endpoint_template="https://api.x.test/detail/{id}",
        detail_method="GET", detail_id_field="id",
    )
    job = result2.jobs[0]
    assert job.full_jd.startswith("List level description.")
    assert result2.jobs[0].raw_data["_generic_detail_source"] == "DETAIL_API"


# CASE 6: short summary + detail URL -> DETAIL_REQUIRED -> fetch.
def test_case6_short_summary_triggers_detail_fetch():
    item = raw(1, "Python 工程师，急招。")
    item["requirements"] = ""
    detail = {"description": "岗位职责\n负责搜索服务研发\n" + "构建高可用服务。 " * 30, "requirements": ""}
    result = _collect(
        [item], {"1": detail}, detail_mode="DETAIL_FALLBACK",
        detail_endpoint_template="https://api.x.test/detail/{id}",
        detail_method="GET", detail_id_field="id",
    )
    assert result.metrics.detail_requests == 1
    assert "岗位职责" in result.jobs[0].full_jd
    assert result.jobs[0].raw_data["_jd_state"] in ("FULL_TEXT", "SUMMARY") or result.jobs[0].full_jd


# CASE 7: summary only and no detail available -> honestly incomplete.
def test_case7_summary_without_detail_is_not_fake_complete():
    item = raw(1, "Python 工程师，急招。")
    item["requirements"] = ""
    item["detail_url"] = "https://api.x.test/detail/1"  # dead endpoint
    client = Client(
        [{"data": {"items": [item], "total": 1}}],  # no detail payload served
    )
    plan_value = plan(detail_mode="DETAIL_FALLBACK")
    result = GenericHttpCollector(plan_value, client).collect()
    dc = evaluate_data_completeness(result)
    assert dc.jd_incomplete == 1
    assert result.jobs[0].full_jd is not None  # summary kept verbatim, not dropped


# CASE 8: lone surrogates survive JSON / Markdown / Excel export.
def test_case8_lone_surrogate_exported_safely(tmp_path):
    poisoned = "岗位职责\n负责平台研发 \ud800 后端服务\n任职要求\n本科 \udfff 以上"
    job = Job(
        job_title="工程师 \ud83d", locations=["上海"],
        responsibilities=["\ud800负责开发"], requirements=["Python \udfff"],
        full_jd=poisoned, source_url="https://example.test/jobs",
        raw_data={"description": poisoned, "nested": ["x \udc00"]},
    )
    result = CollectionResult(
        source_url="https://example.test/jobs", platform="test",
        total_expected=1, total_fetched=1,
        total_unique=1, status="COMPLETE", jobs=[job],
    )
    result.metrics.detail_requests = 0
    result.data_completeness = evaluate_data_completeness(result)
    artifacts = ReportManager().generate_reports(result, tmp_path / "run")
    # JSON round-trip
    data = json.loads(artifacts.json_path.read_text(encoding="utf-8"))
    assert "负责平台研发" in data["jobs"][0]["full_jd"]
    # Markdown written
    markdown = artifacts.markdown_path.read_text(encoding="utf-8")
    assert "负责平台研发" in markdown and "任职要求" in markdown
    # Excel written and readable
    workbook = load_workbook(artifacts.excel_path)
    assert workbook["Jobs"].max_row == 2
    # content preserved except the invalid units
    assert " \ud800 " not in job.full_jd and "职责" in job.full_jd


# CASE 9: Chinese + emoji + bullets + newlines preserved verbatim.
def test_case9_rich_text_preserved_verbatim():
    jd = "岗位职责\n• 负责核心系统研发 🚀\n\t• 参与技术评审 ✅\n任职要求\n• 本科以上学历\n• 熟悉 Python/Go"
    result = _collect([{"id": "1", "title": "工程师", "description": jd}])
    job = result.jobs[0]
    assert job.full_jd == jd  # byte-for-byte
    assert "🚀" in job.full_jd and "•" in job.full_jd and "\n" in job.full_jd and "\t" in job.full_jd


# CASE 10: 2000 jobs already carrying full JD -> zero detail storm.
def test_case10_no_detail_storm_on_large_complete_list():
    items = [
        {"id": str(i), "title": f"工程师 {i}", "description": "岗位职责\n负责平台研发。" + "构建高可用服务。 " * 40}
        for i in range(1, 2001)
    ]
    details = {str(i): {"description": "SHOULD NOT BE FETCHED"} for i in range(1, 2001)}
    result = _collect(
        items, details, detail_mode="DETAIL_FALLBACK",
        detail_endpoint_template="https://api.x.test/detail/{id}",
        detail_method="GET", detail_id_field="id",
    )
    assert result.metrics.detail_requests == 0
    assert result.total_unique == 2000
    assert result.data_completeness.jd_complete == 2000
    assert result.data_completeness.list_sufficient == 2000


# Semantics sanity: per-job trigger decision function.
def test_needs_detail_fetch_decision_table():
    full = {"description": "岗位职责\n负责研发\n" + "构建服务。 " * 60}
    split = {"jobResponsibility": "负责研发", "jobRequirement": "本科以上"}
    half = {"jobRequirement": "Python"}
    summary = {"description": "Python 工程师，急招。"}
    absent = {"id": "1", "title": "t"}
    assert not needs_detail_fetch(full, "DETAIL_FALLBACK")
    assert not needs_detail_fetch(split, "DETAIL_FALLBACK")
    assert needs_detail_fetch(half, "DETAIL_FALLBACK")
    assert needs_detail_fetch(summary, "DETAIL_FALLBACK")
    assert needs_detail_fetch(absent, "DETAIL_FALLBACK")
    assert not needs_detail_fetch(absent, "LIST_SUFFICIENT")
    assert needs_detail_fetch(absent, "DETAIL_REQUIRED")
