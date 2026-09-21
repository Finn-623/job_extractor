"""最小测试：验证 examples/example_output/ 为真实导出管线的合法产物。

覆盖两层：
1. 仓库内已提交的示例文件可直接解析（json / csv / xlsx / md / collection.json）。
2. scripts/generate_example_output.py 的 main() 在 tmp_path 里可重复生成同样的 5 个文件。
"""

import csv
import json
from pathlib import Path

from openpyxl import load_workbook

REPO = Path(__file__).resolve().parents[1]
EXAMPLE_DIR = REPO / "examples" / "example_output"

EXPECTED_FILES = {
    "jobs.json",
    "jobs.csv",
    "jobs.xlsx",
    "report.md",
    "collection.json",
}


def test_committed_files_exist():
    assert EXAMPLE_DIR.is_dir()
    assert EXPECTED_FILES.issubset({p.name for p in EXAMPLE_DIR.iterdir()})


def test_jobs_json_parses_with_fictional_company():
    data = json.loads((EXAMPLE_DIR / "jobs.json").read_text(encoding="utf-8"))
    assert isinstance(data, dict)
    assert len(data["jobs"]) == 5
    for job in data["jobs"]:
        assert job["company"] == "Example Robotics Ltd."
        assert job["detail_url"].startswith("https://jobs.example.com/")
        assert job["full_jd"]


def test_jobs_csv_has_header_and_rows():
    with (EXAMPLE_DIR / "jobs.csv").open(encoding="utf-8-sig", newline="") as f:
        rows = list(csv.reader(f))
    assert len(rows) == 6  # header + 5 jobs
    assert rows[0][:4] == ["id", "title", "company", "location"]


def test_jobs_xlsx_has_three_sheets():
    wb = load_workbook(EXAMPLE_DIR / "jobs.xlsx", read_only=True)
    assert wb.sheetnames == ["Summary", "Jobs", "Data Quality"]
    assert wb["Jobs"].max_row == 6


def test_collection_json_mode_generic():
    data = json.loads((EXAMPLE_DIR / "collection.json").read_text(encoding="utf-8"))
    assert data["collection_mode"] == "generic"
    assert data["source_url"] == "https://example.com/careers"


def test_report_md_mentions_company():
    text = (EXAMPLE_DIR / "report.md").read_text(encoding="utf-8")
    assert "Example Robotics Ltd." in text
    assert "COMPLETE" in text


def test_generator_reproduces_files(tmp_path):
    import importlib.util

    spec = importlib.util.spec_from_file_location(
        "generate_example_output", REPO / "scripts" / "generate_example_output.py"
    )
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)

    module.main(output_dir=tmp_path)
    produced = {p.name for p in tmp_path.iterdir()}
    assert EXPECTED_FILES.issubset(produced)
    regenerated = json.loads((tmp_path / "jobs.json").read_text(encoding="utf-8"))
    assert len(regenerated["jobs"]) == 5
