"""STEP 50 — Export Hardening tests.

Covers: illegal control chars, newline/tab preservation, cell-length
truncation, formula-injection escape, CJK/emoji retention, lone surrogates,
mixed types, row-level fault isolation, sheet-name safety, and a
Chery-like large export with adversarial content.
"""
from __future__ import annotations

import json
import random

import pytest
from openpyxl import load_workbook

from job_extractor.models import CollectionResult, CollectionMetrics, Job
from job_extractor.reporting import ReportManager
from job_extractor.reporting.sanitize import (
    EXCEL_CELL_LIMIT,
    ExcelSanitizeStats,
    TRUNCATION_MARKER,
    sanitize_excel_value,
    safe_sheet_name,
)

try:  # openpyxl raises this on illegal XML chars at cell assignment
    from openpyxl.utils.exceptions import IllegalCharacterError
except ImportError:  # pragma: no cover
    IllegalCharacterError = Exception


def job(i=1, title=None, full="完整 JD", resp=None, req=None, **overrides):
    base = dict(company="测试 & Co", job_id=str(i), job_title=title or f"工程师 {i} 🚀",
        locations=["上海", "北京"],
        responsibilities=resp if resp is not None else ["开发 <系统>", "保障质量"],
        requirements=req if req is not None else ["Python", "沟通能力"],
        full_jd=full, detail_url=f"https://example.test/jobs/{i}",
        apply_url=f"https://example.test/jobs/{i}/apply",
        source_url="https://example.test/jobs")
    base.update(overrides)
    return Job(**base)


def result(jobs=None, warnings=None, errors=None, company="测试公司"):
    jobs = jobs if jobs is not None else [job()]
    r = CollectionResult(source_url="https://example.test/jobs", platform="test",
        company=company, scope={"type": "campus"},
        metrics=CollectionMetrics(list_requests=1, list_pages=1, elapsed_seconds=1.25,
            jd_strategy="LIST_SUFFICIENT"),
        total_expected=len(jobs), total_fetched=len(jobs), total_unique=len(jobs),
        status="COMPLETE", jobs=jobs, warnings=warnings or [], errors=errors or [])
    return r


# 1. Illegal XML control chars are removed; string survives export.
def test_illegal_control_characters_removed(tmp_path):
    stats = ExcelSanitizeStats()
    dirty = "Hello\x00World\x0bTest\x1f"
    clean = sanitize_excel_value(dirty, stats)
    assert clean == "HelloWorldTest"
    assert stats.illegal_chars_removed == 3 and stats.cells_sanitized == 1
    r = result([job(title=dirty, full=dirty)])
    a = ReportManager().generate_reports(r, tmp_path)  # must not raise
    assert load_workbook(a.excel_path)["Jobs"].cell(2, 3).value == "HelloWorldTest"


# 2. Newline and tab are preserved (allowed whitespace controls).
def test_newline_and_tab_preserved():
    text = "line1\nline2\tcolumn\r\nline3"
    out = sanitize_excel_value(text)
    assert "\n" in out and "\t" in out
    assert "\r" not in out  # CRLF normalized to LF


# 3. Cells longer than 32767 chars are truncated with marker, final length ≤ limit.
def test_long_cell_truncated_with_marker():
    stats = ExcelSanitizeStats()
    long = "长" * 40000
    out = sanitize_excel_value(long, stats)
    assert len(out) <= EXCEL_CELL_LIMIT
    assert out.endswith(TRUNCATION_MARKER) and stats.cells_truncated == 1
    # JSON retains the full value (no data loss at the model layer).
    r = result([job(full=long)])
    assert r.jobs[0].full_jd == long


# 4. Formula injection prefixes are escaped with a leading apostrophe.
@pytest.mark.parametrize("payload", ["=SUM(A1:A2)", "+123", "-10+20", "@cmd"])
def test_formula_injection_escaped(payload):
    stats = ExcelSanitizeStats()
    out = sanitize_excel_value(payload, stats)
    assert out == "'" + payload
    assert stats.formula_escaped == 1


# 5. Chinese and emoji content is retained unchanged.
def test_chinese_and_emoji_retained(tmp_path):
    title = "#研发 | AI *工程师* _中文_ 🚀"
    a = ReportManager().generate_reports(result([job(title=title)]), tmp_path)
    cell = load_workbook(a.excel_path)["Jobs"].cell(2, 3).value
    assert cell == title  # '#' is not a formula prefix; content intact
    assert "中文" in a.markdown_path.read_text(encoding="utf-8")
    assert json.loads(a.json_path.read_text(encoding="utf-8"))["jobs"][0]["job_title"] == title


# 6. Lone surrogates do not break workbook.save (Excel path only).
# NOTE: JSON (pydantic UTF-8 serialization) and Markdown (write_text) still
# reject lone surrogates — that is a documented cross-artifact limitation,
# so this test targets ExcelReporter directly.
def test_lone_surrogate_does_not_break_save(tmp_path):
    from job_extractor.reporting.excel_reporter import ExcelReporter
    r = result([job(title="bad\ud800surrogate", full="JD\udffftext")])
    path = ExcelReporter().generate(r, tmp_path / "jobs.xlsx")  # must not raise
    s = load_workbook(path)["Jobs"]
    assert s.cell(2, 3).value == "badsurrogate"
    assert "\ud800" not in s.cell(2, 15).value


# 7. Mixed value types: non-strings pass through untouched.
def test_mixed_value_types_passthrough():
    stats = ExcelSanitizeStats()
    from datetime import datetime
    dt = datetime(2024, 1, 2, 3, 4, 5)
    assert sanitize_excel_value(None, stats) is None
    assert sanitize_excel_value(123, stats) == 123
    assert sanitize_excel_value(1.5, stats) == 1.5
    assert sanitize_excel_value(True, stats) is True
    assert sanitize_excel_value(dt, stats) is dt
    assert sanitize_excel_value(["a", "b"], stats) == ["a", "b"]
    assert sanitize_excel_value({"k": 1}, stats) == {"k": 1}
    assert stats.cells_sanitized == 0


# 8. Row-level fault isolation: 100 rows with one bad cell keeps all rows.
def test_row_fault_isolation_keeps_all_rows(tmp_path, monkeypatch):
    r = result([job(i) for i in range(100)])
    # Force a failure inside the sanitizer path for exactly one cell.
    original_init = ExcelSanitizeStats.__init__

    class FlakyStats(ExcelSanitizeStats):
        def __init__(self):
            super().__init__()
            self.calls = 0

    flaky = FlakyStats()
    real_sanitize = sanitize_excel_value

    import job_extractor.reporting.excel_reporter as er
    real_sanitize = sanitize_excel_value

    def flaky_excel_text(value):
        if isinstance(value, str) and "\x00INJECT\x00" in value:
            raise ValueError("simulated cell failure")
        return real_sanitize(value, None)

    monkeypatch.setattr(er, "excel_text", flaky_excel_text)
    r.jobs[4].job_title = "岗位\x00INJECT\x00五"  # exactly one bad cell
    a = ReportManager().generate_reports(r, tmp_path)
    s = load_workbook(a.excel_path)["Jobs"]
    assert s.max_row - 1 == 100  # no rows lost
    assert s.cell(6, 3).value == "[EXPORT_ERROR]"  # bad cell isolated
    assert s.cell(2, 3).value != "[EXPORT_ERROR]"  # other rows unaffected
    assert any("excel_cell_write_failures=1" in w for w in r.warnings)


# 9. Sheet-name safety for long/illegal company names.
def test_sheet_name_sanitized():
    assert safe_sheet_name("A" * 50) == "A" * 31
    assert safe_sheet_name('a:b\\c/d?e*f[g]h') == "a_b_c_d_e_f_g_h"
    assert safe_sheet_name("") == "Sheet"
    assert safe_sheet_name(None) == "Sheet"
    assert safe_sheet_name("Jobs") == "Jobs"


# 10. Chery-like large export: thousands of rows with control chars,
# over-long JDs and formula prefixes — no rows lost, no IllegalCharacterError.
def test_chery_like_large_export(tmp_path):
    rng = random.Random(50)
    payloads = ["=SUM(A1:A2)", "+bonus", "@risk", "正常文本"]
    jobs = []
    for i in range(2000):
        jd = "职责" * rng.randint(100, 8000)
        if i % 7 == 0:
            jd = jd + "".join(rng.choice("\x00\x0b\x1c") for _ in range(20))
        if i % 11 == 0:
            jd = "长" * 40000
        if i % 13 == 0:
            jd = rng.choice(payloads) + jd
        jobs.append(job(i, title=f"岗位{i}\x0b" if i % 9 == 0 else f"岗位{i}", full=jd))
    r = result(jobs, company="奇瑞汽车")
    a = ReportManager().generate_reports(r, tmp_path)  # must not raise IllegalCharacterError
    s = load_workbook(a.excel_path, read_only=True)["Jobs"]
    assert s.max_row - 1 == 2000  # every row exported
    data = json.loads(a.json_path.read_text(encoding="utf-8"))
    assert len(data["jobs"]) == 2000  # JSON untouched
