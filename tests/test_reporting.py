import json,time
from openpyxl import load_workbook
from job_extractor.models import CollectionResult,CollectionMetrics,Job
from job_extractor.reporting import ReportManager,ReportArtifacts
from job_extractor.reporting.excel_reporter import HEADERS,quality_rows
from job_extractor.reporting.field_formatters import job_completeness,locations_text,numbered,truncate_excel
from job_extractor.reporting.markdown_reporter import MarkdownReporter
from job_extractor.runtime import evaluate_data_completeness

def job(i=1,full="完整 JD",resp=None,req=None,title=None):
    return Job(company="测试 & Co",job_id=str(i),job_title=title or f"工程师 {i} 🚀",locations=["上海","北京"],
        responsibilities=resp if resp is not None else ["开发 <系统>","保障质量"],requirements=req if req is not None else ["Python","沟通能力"],
        full_jd=full,detail_url=f"https://example.test/jobs/{i}",apply_url=f"https://example.test/jobs/{i}/apply",source_url="https://example.test/jobs")
def result(jobs=None,warnings=None,errors=None):
    jobs=jobs if jobs is not None else [job()]
    r=CollectionResult(source_url="https://example.test/jobs",platform="test",company="测试公司",scope={"type":"campus"},
        metrics=CollectionMetrics(list_requests=1,list_pages=1,elapsed_seconds=1.25,jd_strategy="LIST_SUFFICIENT"),
        total_expected=len(jobs),total_fetched=len(jobs),total_unique=len(jobs),status="COMPLETE",jobs=jobs,warnings=warnings or [],errors=errors or [])
    r.data_completeness=evaluate_data_completeness(r);return r
def test_report_manager_and_artifacts(tmp_path):
    a=ReportManager().generate_reports(result(),tmp_path/"safe_run");assert isinstance(a,ReportArtifacts) and a.json_path.exists() and a.excel_path.exists() and a.markdown_path.exists()
def test_excel_sheet_names_summary_and_rows(tmp_path):
    a=ReportManager().generate_reports(result([job(1),job(2)]),tmp_path);w=load_workbook(a.excel_path)
    assert w.sheetnames==["Summary","Jobs","Data Quality"] and w["Jobs"].max_row==3
    values={w["Summary"].cell(i,1).value:w["Summary"].cell(i,2).value for i in range(1,w["Summary"].max_row+1)}
    assert values["Company"]=="测试公司" and values["Unique"]=="2"
def test_jobs_header_order_filter_freeze(tmp_path):
    a=ReportManager().generate_reports(result(),tmp_path);s=load_workbook(a.excel_path)["Jobs"]
    assert [x.value for x in s[1]]==HEADERS and s.freeze_panes=="A2" and s.auto_filter.ref
def test_locations_and_numbered_formatting(): assert locations_text(["上海","北京"])=="上海；北京" and numbered(["A","B"])=="1. A\n2. B"
def test_completeness_formatting():
    assert job_completeness(job())=="COMPLETE"
    assert job_completeness(job(full="",resp=[],req=[]))=="MISSING_JD"
    assert job_completeness(job(req=[]))=="STRUCTURE_PARTIAL"
def test_quality_rows_and_no_issue_case(tmp_path):
    assert not quality_rows(result())
    a=ReportManager().generate_reports(result(),tmp_path);assert load_workbook(a.excel_path)["Data Quality"]["A2"].value=="No data quality issues."
def test_quality_warning_and_error_are_distinct():
    rows=quality_rows(result([job(full="",resp=[],req=[])],warnings=["source warning"],errors=["timeout"]))
    assert any(x[2]=="SOURCE DATA WARNING" for x in rows) and any(x[2]=="COLLECTION ERROR" for x in rows)
def test_urls_are_hyperlinks(tmp_path):
    s=load_workbook(ReportManager().generate_reports(result(),tmp_path).excel_path)["Jobs"]
    assert s.cell(2,16).hyperlink.target.endswith("/1") and s.cell(2,17).hyperlink.target.endswith("/apply")
def test_markdown_auto_full_for_small_result(tmp_path):
    path=tmp_path/"small.md";MarkdownReporter().generate(result(),path);assert "#### Full JD" in path.read_text(encoding="utf-8")
def test_markdown_auto_compact_over_50(tmp_path):
    path=tmp_path/"large.md";MarkdownReporter().generate(result([job(i) for i in range(51)]),path);text=path.read_text(encoding="utf-8")
    assert "available in jobs.xlsx and jobs.json" in text and "#### Full JD" not in text
def test_markdown_explicit_modes(tmp_path):
    p=tmp_path/"x.md";MarkdownReporter().generate(result([job(i) for i in range(51)]),p,"full");assert "#### Full JD" in p.read_text(encoding="utf-8")
def test_unicode_and_special_characters(tmp_path):
    r=result([job(title="#研发 | AI *工程师* _中文_ 🚀")]);a=ReportManager().generate_reports(r,tmp_path)
    assert "中文" in a.markdown_path.read_text(encoding="utf-8") and load_workbook(a.excel_path)["Jobs"].cell(2,3).value.startswith("#研发")
    assert json.loads(a.json_path.read_text(encoding="utf-8"))["jobs"][0]["job_title"].endswith("🚀")
def test_empty_optional_fields(tmp_path):
    j=Job(job_title="Minimal",source_url="https://x.test");a=ReportManager().generate_reports(result([j]),tmp_path)
    assert load_workbook(a.excel_path)["Jobs"].max_row==2
def test_long_jd_is_truncated_only_in_excel(tmp_path):
    long="长"*40000;r=result([job(full=long)]);a=ReportManager().generate_reports(r,tmp_path)
    cell=load_workbook(a.excel_path)["Jobs"].cell(2,15).value
    assert len(cell)<=32767 and "EXCEL_CELL_TRUNCATED" in cell and json.loads(a.json_path.read_text(encoding="utf-8"))["jobs"][0]["full_jd"]==long
def test_1000_jobs_and_same_job_count(tmp_path):
    r=result([job(i) for i in range(1000)]);started=time.perf_counter();a=ReportManager().generate_reports(r,tmp_path);elapsed=time.perf_counter()-started
    assert load_workbook(a.excel_path,read_only=True)["Jobs"].max_row-1==1000
    assert len(json.loads(a.json_path.read_text(encoding="utf-8"))["jobs"])==1000
    assert a.markdown_path.read_text(encoding="utf-8").count("\n### ")==1000 and elapsed<20 and a.excel_path.stat().st_size<5_000_000
