from __future__ import annotations
from __future__ import annotations
import logging
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Alignment,Border,Font,PatternFill,Side
from job_extractor.models import CollectionResult,Job
from job_extractor.reporting.field_formatters import excel_text,job_completeness,locations_text,numbered,set_excel_stats,scope_text,truncate_excel
from job_extractor.reporting.sanitize import ExcelSanitizeStats,safe_sheet_name

logger=logging.getLogger(__name__)

HEADERS=["序号","公司","岗位名称","Job ID","招聘类型","岗位类别","部门","工作地点","学历","专业","招聘人数","发布时间","岗位职责","任职要求","完整 JD","Detail URL","Apply URL","Source URL","数据完整性"]
WIDTHS=[8,20,35,25,15,20,25,25,18,30,12,20,55,55,70,45,45,45,22]
HEADER_FILL=PatternFill("solid",fgColor="1F4E78");SECTION_FILL=PatternFill("solid",fgColor="D9EAF7");THIN=Side(style="thin",color="B7C9D6")
CELL_PLACEHOLDER="[EXPORT_ERROR]"

def quality_rows(result:CollectionResult):
    rows=[]
    for job in result.jobs:
        status=job_completeness(job)
        if status=="MISSING_JD":rows.append((job.job_id,job.job_title,"SOURCE_MISSING_JD","No full JD was available from the collected official source.",job.detail_url))
        elif status=="STRUCTURE_PARTIAL":
            missing=[]
            if not job.responsibilities:missing.append("responsibilities")
            if not job.requirements:missing.append("requirements")
            rows.append((job.job_id,job.job_title,"STRUCTURED_FIELDS_UNAVAILABLE",f"Full JD preserved; structured {', '.join(missing)} unavailable.",job.detail_url))
        if len(job.full_jd or "")>32767:rows.append((job.job_id,job.job_title,"EXCEL_CELL_TRUNCATED","Excel cell truncated; jobs.json retains the full value.",job.detail_url))
    for warning in result.warnings:rows.append((None,None,"SOURCE DATA WARNING",warning,None))
    for error in result.errors:rows.append((None,None,"COLLECTION ERROR",error,None))
    return rows

class ExcelReporter:
    def generate(self,result:CollectionResult,path:Path)->Path:
        stats=ExcelSanitizeStats();set_excel_stats(stats)
        try:
            workbook=Workbook();summary=workbook.active;summary.title="Summary";jobs=workbook.create_sheet(safe_sheet_name("Jobs"));quality=workbook.create_sheet(safe_sheet_name("Data Quality"))
        finally:
            set_excel_stats(None)
        sections=[("COLLECTION",[("Company",result.company),("Platform",result.platform),("Source URL",result.source_url),("Recruitment scope",scope_text(result.scope)),("Collection time",result.finished_at or result.started_at),("Collection status",result.status)]),
            ("JOB COUNTS",[("Expected",result.total_expected),("Fetched",result.total_fetched),("Unique",result.total_unique)]),
            ("DATA QUALITY",[("Complete jobs",result.data_completeness.complete_jobs),("Missing JD",result.data_completeness.missing_jd_jobs),("Missing responsibilities",result.data_completeness.missing_responsibilities_jobs),("Missing requirements",result.data_completeness.missing_requirements_jobs),("Completeness ratio",result.data_completeness.completeness_ratio)]),
            ("PERFORMANCE",[("Elapsed",result.metrics.elapsed_seconds),("Pages",result.metrics.list_pages),("List requests",result.metrics.list_requests),("Detail requests",result.metrics.detail_requests),("JD strategy",result.metrics.jd_strategy)]),
            ("WARNINGS",[("Warnings count",len(result.warnings)),("Errors count",len(result.errors))])]+stats_section(stats)
        row=1
        for name,items in sections:
            summary.cell(row,1,name);summary.cell(row,1).font=Font(bold=True);summary.cell(row,1).fill=SECTION_FILL;summary.merge_cells(start_row=row,start_column=1,end_row=row,end_column=2);row+=1
            for metric,value in items:summary.cell(row,1,metric);summary.cell(row,2,_safe_cell(value,stats));row+=1
            row+=1
        summary.column_dimensions["A"].width=28;summary.column_dimensions["B"].width=70
        jobs.append(HEADERS)
        for cell in jobs[1]:cell.font=Font(bold=True,color="FFFFFF");cell.fill=HEADER_FILL;cell.alignment=Alignment(horizontal="center");cell.border=Border(bottom=THIN)
        for index,job in enumerate(result.jobs,1):
            resp,_=truncate_excel(numbered(job.responsibilities));req,_=truncate_excel(numbered(job.requirements));full,_=truncate_excel(job.full_jd)
            row=[index,job.company,job.job_title,job.job_id,job.recruitment_type,job.job_category,job.department,locations_text(job.locations),job.education,job.major,job.headcount,job.publish_date,resp,req,full,"查看岗位" if job.detail_url else None,"前往投递" if job.apply_url else None,"来源网页" if job.source_url else None,job_completeness(job)]
            jobs.append([_safe_cell(x,stats) for x in row]);r=jobs.max_row;jobs.row_dimensions[r].height=45
            for col,url in ((16,job.detail_url),(17,job.apply_url),(18,job.source_url)):
                if url:jobs.cell(r,col).hyperlink=url;jobs.cell(r,col).style="Hyperlink"
            for col in range(1,len(HEADERS)+1):jobs.cell(r,col).alignment=Alignment(vertical="top",wrap_text=col in (13,14,15))
        jobs.freeze_panes="A2";jobs.auto_filter.ref=jobs.dimensions
        for i,width in enumerate(WIDTHS,1):jobs.column_dimensions[jobs.cell(1,i).column_letter].width=width
        quality.append(["Job ID","Job Title","Issue Type","Description","Detail URL"])
        for cell in quality[1]:cell.font=Font(bold=True,color="FFFFFF");cell.fill=HEADER_FILL
        issues=quality_rows(result)
        if issues:
            for issue in issues:quality.append([_safe_cell(x,stats) for x in issue])
        else:quality.append(["No data quality issues."])
        quality.freeze_panes="A2";quality.auto_filter.ref=quality.dimensions
        for letter,width in zip(("A","B","C","D","E"),(25,35,35,70,45)):quality.column_dimensions[letter].width=width
        _log_export_diagnostics(stats,len(result.jobs),result)
        path.parent.mkdir(parents=True,exist_ok=True);workbook.save(path);return path

def stats_section(stats:ExcelSanitizeStats)->list[tuple[str,list[tuple[str,Any]]]]:
    return [("EXPORT DIAGNOSTICS",[(k.replace("excel_","").replace("_"," ").title(),v) for k,v in stats.to_dict().items() if k!="excel_cell_write_failures"]+[("Cell write failures",stats.cell_write_failures)])]

def _safe_cell(value,stats:ExcelSanitizeStats):
    """Fault-isolated cell write: one bad cell must not kill the export."""
    try:
        return excel_text(value)
    except Exception as exc:  # noqa: BLE001 - recorded, never silent
        stats.cell_write_failures+=1
        logger.warning("excel cell write failed: %s: %s",type(exc).__name__,exc)
        return CELL_PLACEHOLDER

def _log_export_diagnostics(stats:ExcelSanitizeStats,job_count:int,result:CollectionResult)->None:
    diag=stats.to_dict()
    logger.info("excel export diagnostics: jobs=%d %s",job_count,diag)
    if stats.cell_write_failures:
        result.warnings.append(f"excel_cell_write_failures={stats.cell_write_failures}")
